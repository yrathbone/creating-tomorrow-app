"""Resume layouts: Classic (the original look, unchanged), Modern and Traditional.

Every layout holds exactly the same words; only the look changes. Unknown layout names quietly mean
Classic, ATS mode still works in every layout, and the download routes accept the choice.

Run from the backend folder:  py -3.12 -m unittest discover -s tests -t .
"""
import io
import inspect
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from fastapi.testclient import TestClient

import resume_builder as rb

SAMPLE = {
    "name": "Ada Lovelace",
    "contact": "London, UK | 555-0100 | ada@example.com",
    "headline": "Customer Success Leader",
    "summary": ["Led renewals for 40 accounts.", "Built an onboarding playbook."],
    "skills": ["Consultative selling", "Renewals", "Onboarding", "Customer retention"],
    "experience": [
        {"title": "Customer Success Manager", "subtitle": "Northwind Software - 03/21 - Present", "bullets": ["Owned 40 accounts."]},
        {"title": "Account Manager", "subtitle": "Contoso Retail - 06/17 - 02/21", "bullets": ["Managed 120 accounts."]},
    ],
    "education": ["M.S., Information Technology - Illinois Institute of Technology (05/27)"],
    "certifications": ["Certified Talent Professional (CTP)"],
    "additional_sections": [{"heading": "LANGUAGES", "items": ["English (native)", "Spanish (professional)"]}],
}
HEADINGS = {"PROFESSIONAL SUMMARY", "CORE SKILLS & EXPERTISE", "PROFESSIONAL EXPERIENCE", "EDUCATION", "CERTIFICATIONS", "LANGUAGES"}


def build(template=None, ats=False, data=SAMPLE):
    kwargs = {"ats_mode": ats}
    if template is not None:
        kwargs["template"] = template
    return rb.build_resume_bytes(data, **kwargs)


def doc_of(blob):
    return Document(io.BytesIO(blob))


def heading_paragraphs(doc):
    return [p for p in doc.paragraphs if p.text in HEADINGS]


def ppr_children(paragraph):
    pPr = paragraph._p.pPr
    return [child.tag.split("}")[1] for child in pPr] if pPr is not None else []


class LayoutLooks(unittest.TestCase):
    def test_classic_is_the_default_and_keeps_the_original_look(self):
        doc = doc_of(build())
        name = doc.paragraphs[0]
        self.assertEqual(name.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(name.runs[0].font.name, "Calibri")
        self.assertEqual(name.runs[0].font.size, Pt(11))
        for heading in heading_paragraphs(doc):
            xml = heading._p.xml
            self.assertIn('w:fill="CADCF2"', xml, heading.text)  # the shaded bar
            self.assertNotIn("w:pBdr", xml)
            self.assertEqual(heading.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_unknown_layout_names_mean_classic(self):
        classic = doc_of(build("classic")).element.xml
        for bad in ("fancy", "", "MODERN ", None, 7):
            self.assertEqual(doc_of(rb.build_resume_bytes(SAMPLE, template=bad)).element.xml, classic, repr(bad))

    def test_modern_is_left_aligned_navy_with_a_line_under_each_title(self):
        doc = doc_of(build("modern"))
        name = doc.paragraphs[0]
        self.assertEqual(name.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        self.assertEqual(name.runs[0].font.size, Pt(22))
        self.assertEqual(str(name.runs[0].font.color.rgb), "1F3A8A")
        self.assertEqual(str(name.runs[2].font.color.rgb), "4B5565")  # the contact line is softer
        headings = heading_paragraphs(doc)
        self.assertEqual({h.text for h in headings}, HEADINGS)
        for h in headings:
            self.assertIn('<w:bottom w:val="single"', h._p.xml)
            self.assertIn('w:color="1F3A8A"', h._p.xml)
            self.assertNotIn("w:shd", h._p.xml)
            self.assertEqual(h.alignment, WD_ALIGN_PARAGRAPH.LEFT)

    def test_traditional_is_serif_with_black_lines(self):
        doc = doc_of(build("traditional"))
        self.assertEqual(doc.paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(doc.paragraphs[0].runs[0].font.size, Pt(18))
        fonts = {run.font.name for p in doc.paragraphs for run in p.runs if run.text.strip()}
        self.assertEqual(fonts, {"Times New Roman"})  # every word is set in the serif font
        self.assertEqual(doc.styles["Normal"].font.name, "Times New Roman")  # and so is the document default
        for h in heading_paragraphs(doc):
            self.assertIn('w:color="000000"', h._p.xml)
            self.assertNotIn("w:shd", h._p.xml)

    def test_heading_lines_are_placed_where_word_expects_them(self):
        # Word is strict about the order of paragraph properties: the border must come before spacing and alignment
        for template in ("modern", "traditional"):
            for h in heading_paragraphs(doc_of(build(template))):
                kids = ppr_children(h)
                self.assertIn("pBdr", kids)
                for later in ("spacing", "jc"):
                    if later in kids:
                        self.assertLess(kids.index("pBdr"), kids.index(later), (template, h.text, kids))


class SameWordsEverywhere(unittest.TestCase):
    def texts(self, template, ats=False):
        return [p.text.strip() for p in doc_of(build(template, ats)).paragraphs if p.text.strip()]

    def test_every_layout_has_exactly_the_same_words_in_the_same_order(self):
        base = self.texts("classic")
        for template in ("modern", "traditional"):
            self.assertEqual(self.texts(template), base, template)

    def test_ats_mode_still_makes_the_skills_a_single_column_in_every_layout(self):
        def column_counts(template, ats):
            doc = doc_of(build(template, ats))
            return [s._sectPr.find(qn("w:cols")).get(qn("w:num")) for s in doc.sections if s._sectPr.find(qn("w:cols")) is not None]

        for template in ("classic", "modern", "traditional"):
            self.assertIn("2", column_counts(template, False), template)
            self.assertNotIn("2", column_counts(template, True), template)

    def test_empty_sections_are_still_skipped_in_every_layout(self):
        sparse = dict(SAMPLE, skills=[], experience=[], education=[], certifications=[], additional_sections=[], summary=None)
        for template in ("classic", "modern", "traditional"):
            headings = {p.text for p in doc_of(build(template, data=sparse)).paragraphs} & HEADINGS
            self.assertEqual(headings, set(), template)


class DownloadRoutes(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import main
        cls.client = TestClient(main.app)

    def generate(self, **extra):
        body = {"resume_data": SAMPLE}
        body.update(extra)
        return self.client.post("/api/generate", json=body)

    def test_generate_accepts_each_layout_and_returns_a_word_file(self):
        sizes = {}
        for template in ("classic", "modern", "traditional"):
            res = self.generate(template=template)
            self.assertEqual(res.status_code, 200, template)
            self.assertIn("wordprocessingml", res.headers["content-type"])
            sizes[template] = doc_of(res.content).paragraphs[0].runs[0].font.size
        self.assertEqual(sizes, {"classic": Pt(11), "modern": Pt(22), "traditional": Pt(18)})

    def test_generate_without_a_layout_or_with_a_wrong_one_gives_classic(self):
        classic = doc_of(self.generate(template="classic").content).element.xml
        self.assertEqual(doc_of(self.generate().content).element.xml, classic)
        self.assertEqual(doc_of(self.generate(template="nonsense").content).element.xml, classic)

    def test_the_career_download_route_takes_a_layout_too(self):
        from career.routes import download_resume_version
        params = inspect.signature(download_resume_version).parameters
        self.assertIn("template", params)
        # no explicit choice means "the person's saved layout" (see test_resume_layout.py)
        self.assertIsNone(params["template"].default)


if __name__ == "__main__":
    unittest.main()
