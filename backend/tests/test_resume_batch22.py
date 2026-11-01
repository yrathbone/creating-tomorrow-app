"""Phase 02, Batch 2.2: one-column skills by default, the name and the contact line as two paragraphs, one date format,
and a Career Profile that is never reduced by what a resume chooses to show.

Run from the backend folder:  py -3.12 -m unittest discover -s tests -t .
"""
import io
import os
import re
import unittest
from unittest import mock

from docx import Document
from docx.oxml.ns import qn

import abuse_guard
import resume_builder as rb
from tests.test_ownership import OwnershipTestCase

LAYOUTS = ("classic", "modern", "traditional")
MANY = [f"Skill {i:02d}" for i in range(1, 61)]
BASE = {
    "name": "Casey Rowan", "contact": "Naperville, IL | casey.rowan@example.com | (555) 010-0166",
    "summary": "Operations analyst.", "skills": MANY,
    "experience": [{"title": "Analyst", "subtitle": "Brightline Services, Naperville, IL — 02/18 – present", "bullets": ["Built the weekly dashboard"]}],
    "education": ["B.S., Business Analytics — NIU, 2017"], "certifications": [],
}


def build(data=BASE, template="classic", ats=False):
    return Document(io.BytesIO(rb.build_resume_bytes(data, ats_mode=ats, template=template)))


def column_counts(doc):
    return [s._sectPr.find(qn("w:cols")).get(qn("w:num")) for s in doc.sections if s._sectPr.find(qn("w:cols")) is not None]


def skills_text(doc):
    paragraphs = [p.text for p in doc.paragraphs]
    start = paragraphs.index("CORE SKILLS & EXPERTISE") + 1
    end = paragraphs.index("PROFESSIONAL EXPERIENCE")
    return paragraphs[start:end]


class OneColumnByDefault(unittest.TestCase):
    def test_a_resume_with_no_skills_setting_is_one_column_in_every_layout(self):
        for template in LAYOUTS:
            with self.subTest(template=template):
                doc = build(template=template)
                self.assertEqual(len(doc.sections), 1)
                self.assertNotIn("2", column_counts(doc))

    def test_all_sixty_skills_are_kept_in_order_and_take_one_line_each_run(self):
        for template in LAYOUTS:
            with self.subTest(template=template):
                lines = skills_text(build(template=template))
                self.assertEqual(lines, [", ".join(MANY)])  # nothing dropped, reworded or reordered

    def test_labelled_skills_stay_as_labelled_lines_and_plain_ones_join_between_them(self):
        data = dict(BASE, skills=["Data: Excel, SQL", "Python", "R", "People: Coaching", "Hiring"])
        self.assertEqual(skills_text(build(data)), ["Data: Excel, SQL", "Python, R", "People: Coaching", "Hiring"])

    def test_skills_that_contain_commas_are_never_merged_into_one_run(self):
        data = dict(BASE, skills=["Budgeting, forecasting", "Hiring", "Coaching"])
        self.assertEqual(skills_text(build(data)), ["Budgeting, forecasting", "Hiring", "Coaching"])

    def test_two_columns_are_still_available_when_asked_for(self):
        doc = build(dict(BASE, skills_style="columns", skills=MANY[:6]))
        self.assertIn("2", column_counts(doc))
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == "List Bullet"][:6], MANY[:6])

    def test_the_grouped_style_and_ats_mode_behave_as_before(self):
        grouped = skills_text(build(dict(BASE, skills_style="grouped", skills=["A: x, y", "B: z"])))
        self.assertEqual(grouped, ["A: x, y", "B: z"])
        ats = build(dict(BASE, skills=MANY[:3]), ats=True)
        self.assertEqual([p.text for p in ats.paragraphs if p.style.name == "List Bullet"][:3], MANY[:3])
        self.assertNotIn("2", column_counts(ats))

    def test_the_default_layout_is_much_shorter_than_sixty_bullets(self):
        self.assertLess(len(skills_text(build())), 3)

    def test_an_unknown_skills_style_means_the_default(self):
        self.assertEqual(skills_text(build(dict(BASE, skills_style="wild"))), [", ".join(MANY)])


class NameAndContact(unittest.TestCase):
    def test_the_name_and_the_contact_line_are_two_paragraphs(self):
        for template in LAYOUTS:
            with self.subTest(template=template):
                paragraphs = [p.text for p in build(template=template).paragraphs]
                self.assertEqual(paragraphs[:2], ["Casey Rowan", "Naperville, IL | casey.rowan@example.com | (555) 010-0166"])
                for text in paragraphs[:2]:
                    self.assertNotIn("\n", text)

    def test_the_name_paragraph_stays_with_the_contact_line_and_is_bold(self):
        first = build().paragraphs[0]
        self.assertTrue(first.paragraph_format.keep_with_next)
        self.assertTrue(first.runs[0].bold)

    def test_a_missing_contact_or_name_leaves_no_empty_paragraph(self):
        for data in (dict(BASE, contact=""), dict(BASE, name=""), dict(BASE, name="", contact="")):
            with self.subTest(name=data["name"], contact=data["contact"]):
                self.assertEqual([i for i, p in enumerate(build(data).paragraphs) if not p.text.strip()], [])

    def test_the_words_are_the_same_in_every_layout(self):
        base = [p.text for p in build(template="classic").paragraphs]
        for template in ("modern", "traditional"):
            self.assertEqual([p.text for p in build(template=template).paragraphs], base)


class EmploymentDates(unittest.TestCase):
    CASES = {
        "Brightline Services, Naperville, IL — 02/18 – Present": "Brightline Services, Naperville, IL — Feb 2018 – Present",
        "Acme — Jan 2018 - present": "Acme — Jan 2018 – Present",
        "Acme — January 2018 to June 2020": "Acme — Jan 2018 – Jun 2020",
        "Acme — 03/2015 – 11/2017": "Acme — Mar 2015 – Nov 2017",
        "Acme — 1/99 - 12/04": "Acme — Jan 1999 – Dec 2004",
        "Acme — 2018-2020": "Acme — 2018 – 2020",
        "Acme — 2018-02 – 2019-03": "Acme — Feb 2018 – Mar 2019",
        "Acme — Sept. 2019 – Current": "Acme — Sep 2019 – Present",
        "Acme — 2020 until now": "Acme — 2020 – Present",
        "Feb 2018 – Present": "Feb 2018 – Present",
    }
    UNCHANGED = [
        "Acme — Spring 2025", "Acme — 2019", "Acme, Town", "", "Acme — 13/18 – 05/19", "Acme — Present – 2020",
        "Acme — Summer 2019 – Fall 2020", "Acme — 02/18 – somewhere", "Acme — about 2018 – 2020 (approx.)",
    ]

    def test_ranges_become_one_format(self):
        for raw, expected in self.CASES.items():
            with self.subTest(raw=raw):
                self.assertEqual(rb.standardize_dates(raw), expected)

    def test_anything_not_fully_understood_is_left_exactly_as_written(self):
        for raw in self.UNCHANGED:
            with self.subTest(raw=raw):
                self.assertEqual(rb.standardize_dates(raw), raw)

    def test_no_month_is_ever_invented_for_a_year_only_date(self):
        self.assertEqual(rb.standardize_dates("Acme — 2018 – 2020"), "Acme — 2018 – 2020")
        self.assertNotRegex(rb.standardize_dates("Acme — 2018-2020"), r"[A-Z][a-z]{2} 20")

    def test_running_it_twice_changes_nothing_more(self):
        for raw in self.CASES:
            once = rb.standardize_dates(raw)
            self.assertEqual(rb.standardize_dates(once), once)

    def test_the_document_shows_the_standard_dates_in_every_layout(self):
        for template in LAYOUTS:
            with self.subTest(template=template):
                text = "\n".join(p.text for p in build(template=template).paragraphs)
                self.assertIn("Brightline Services, Naperville, IL — Feb 2018 – Present", text)
                self.assertNotIn("02/18", text)

    def test_only_job_dates_change_not_bullets_education_or_other_text(self):
        data = dict(BASE, summary="Worked 02/18 - 03/19 on a long project.",
                    experience=[{"title": "T", "subtitle": "S — 02/18 – 03/19", "bullets": ["Ran the 02/18 - 03/19 migration"]}],
                    education=["Diploma — 06/20 – 06/22"])
        text = "\n".join(p.text for p in build(data).paragraphs)
        self.assertIn("S — Feb 2018 – Mar 2019", text)
        self.assertIn("Worked 02/18 - 03/19 on a long project.", text)
        self.assertIn("Ran the 02/18 - 03/19 migration", text)
        self.assertIn("Diploma — 06/20 – 06/22", text)


class ProfileSkillsAreNeverReduced(OwnershipTestCase):
    """A resume may show a selection of skills; the Career Profile keeps every one of them."""

    def setUp(self):
        super().setUp()
        for patcher in (
            mock.patch.dict(os.environ, {"RATE_LIMIT_USER": "1000/60"}),
            mock.patch.object(abuse_guard, "LIMITER", abuse_guard.SlidingWindowLimiter()),
        ):
            patcher.start()
            self.addCleanup(patcher.stop)
        abuse_guard.IN_FLIGHT.reset()

    def profile_skills(self):
        return sorted(s["name"] for s in self.call("GET", "/skills", self.a).json())

    def test_building_a_general_resume_with_fewer_skills_leaves_every_profile_skill(self):
        self.onboard(self.a)
        names = [f"Verified skill {i}" for i in range(1, 13)]
        for name in names:
            self.assertEqual(self.call("POST", "/skills", self.a, json={"name": name}).status_code, 200)
        before = self.profile_skills()
        self.assertTrue(set(names) <= set(before))
        selection = {"name": "N", "contact": "c", "summary": "", "skills": ["Verified skill 1, Verified skill 2", "Not in profile at all"],
                     "experience": [], "education": []}
        with mock.patch("career.routes.build_general_resume", return_value={"resume_data": dict(selection)}):
            res = self.call("POST", "/build-general-resume", self.a, json={"name": "N", "contact": "c@example.com"})
        self.assertEqual(res.status_code, 200, res.text)
        self.assertEqual(self.profile_skills(), before)  # nothing removed, added or renamed

    def test_downloading_the_resume_does_not_touch_the_profile_either(self):
        self.onboard(self.a)
        self.call("POST", "/skills", self.a, json={"name": "Only skill"})
        before = self.profile_skills()
        import main
        from fastapi.testclient import TestClient
        res = TestClient(main.app).post("/api/generate", json={"resume_data": dict(BASE, skills=["Something else"]), "ats_mode": False})
        self.assertIn(res.status_code, (200, 401, 403, 422))
        self.assertEqual(self.profile_skills(), before)


if __name__ == "__main__":
    unittest.main()
