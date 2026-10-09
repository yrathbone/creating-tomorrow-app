"""The readable Word copy of a saved Career Profile: formats what is saved, adds and invents nothing."""
import io
import unittest

from docx import Document

from profile_copy import build_profile_copy_bytes, pretty_date

SAMPLE = {
    "exported_at": "2026-10-09T01:19:44.808945+00:00",
    "experiences": [
        {"title": "Treasury Management Officer", "organization": "JPMorgan Chase", "location": None, "start_date": "06/26", "end_date": "Present",
         "description": "Cultivate client relationships.\nLead executive discovery sessions.\n\nNegotiated contract terms", "verified": True},
        {"title": "Analyst", "organization": "Acme", "location": "Chicago, IL", "start_date": "01/20", "end_date": "12/21", "description": "", "verified": False},
    ],
    "education": [
        {"institution": "Illinois Institute of Technology", "degree": "M.S.", "field_of_study": "Applied Information Technology", "graduation_date": "In Progress, 2028"},
        {"institution": "University of Maryland Global Campus", "degree": "Accounting Foundation Certificate", "field_of_study": None, "graduation_date": "2023"},
    ],
    "certifications": [
        {"name": "Certified Treasury Professional (CTP)", "issuer": None, "date": "In Progress"},
        {"name": "Google Data Analytics", "issuer": "Google", "date": None},
    ],
    "skills": [{"name": "Treasury"}, {"name": "api design"}, {"name": "Treasury"}, {"name": "Analytics"}],
    "languages": [{"name": "Spanish", "proficiency": "Fluent"}],
    "applications": [{"job_title": "Product Lead", "company": "Contoso", "applied_on": "2026-10-01", "status": "Applied", "notes": "Referred by a friend"}],
    "scan_history": [{"id": 7, "scan_type": "job_comparison", "job_title": "Product Lead", "summary_text": "Strong fit.\nTwo gaps.", "created_at": "2026-10-02T15:00:00+00:00"}],
    "resume_versions": [{"id": 3, "scan_history_id": 7, "resume_data": {"name": "Yovana Rathbone"}, "created_at": "2026-10-03T15:00:00+00:00"}],
}


def text_of(docx_bytes):
    doc = Document(io.BytesIO(docx_bytes))
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    return [p.text for p in doc.paragraphs], cells


class ReadableCopyTests(unittest.TestCase):
    def setUp(self):
        self.paragraphs, self.cells = text_of(build_profile_copy_bytes(SAMPLE, "Yovana Rathbone", "Chicago, IL | 555-0100 | me@example.com"))
        self.text = "\n".join(self.paragraphs)

    def test_it_is_a_word_document_not_data(self):
        raw = build_profile_copy_bytes(SAMPLE, "Yovana Rathbone", None)
        self.assertEqual(raw[:2], b"PK")  # a .docx is a zip
        self.assertNotIn('"exported_at"', self.text)
        self.assertNotIn("created_at", self.text)

    def test_name_contact_and_date_are_at_the_top(self):
        self.assertEqual(self.paragraphs[0], "Yovana Rathbone")
        self.assertIn("Chicago, IL | 555-0100 | me@example.com", self.text)
        self.assertIn("as of Oct 9, 2026", self.text)

    def test_sections_are_there_in_plain_words(self):
        for heading in ("WORK EXPERIENCE", "EDUCATION", "CERTIFICATIONS", "SKILLS (3)", "LANGUAGES", "APPLICATIONS YOU ARE TRACKING", "SCANS YOU HAVE RUN", "RESUMES YOU HAVE BUILT"):
            self.assertIn(heading, self.paragraphs)

    def test_each_line_of_a_role_description_becomes_its_own_bullet(self):
        self.assertIn("Treasury Management Officer at JPMorgan Chase", self.text)
        self.assertIn("06/26 to Present", self.text)
        for line in ("Cultivate client relationships.", "Lead executive discovery sessions.", "Negotiated contract terms"):
            self.assertIn(line, self.paragraphs)
        self.assertEqual(self.paragraphs.count("Lead executive discovery sessions."), 1)

    def test_unconfirmed_roles_are_marked_and_confirmed_ones_are_not(self):
        self.assertIn("Chicago, IL · 01/20 to 12/21 · not yet confirmed by you", self.paragraphs)
        self.assertNotIn("not yet confirmed by you", [p for p in self.paragraphs if "06/26" in p][0])

    def test_education_and_certifications_skip_empty_parts(self):
        self.assertIn("M.S., Applied Information Technology · Illinois Institute of Technology · In Progress, 2028", self.paragraphs)
        self.assertIn("Accounting Foundation Certificate · University of Maryland Global Campus · 2023", self.paragraphs)
        self.assertIn("Certified Treasury Professional (CTP) · In Progress", self.paragraphs)
        self.assertIn("Google Data Analytics · Google", self.paragraphs)

    def test_skills_are_one_sorted_list_without_repeats(self):
        self.assertIn("Analytics, api design, Treasury", self.paragraphs)

    def test_applications_table_and_history_are_readable(self):
        self.assertIn("Contoso", self.cells)
        self.assertIn("Referred by a friend", self.cells)
        self.assertIn("Job comparison: Product Lead", self.paragraphs)
        self.assertIn("Oct 2, 2026", self.paragraphs)
        self.assertIn("Strong fit.", self.paragraphs)
        self.assertIn("Oct 3, 2026 · for Product Lead · Yovana Rathbone", self.paragraphs)

    def test_an_empty_profile_says_so_and_never_crashes(self):
        paragraphs, cells = text_of(build_profile_copy_bytes({"exported_at": "2026-10-09T00:00:00+00:00"}, None, None))
        self.assertEqual(paragraphs[0], "My Career Profile")
        for what in ("roles", "education", "certifications", "skills", "applications", "scans", "resumes"):
            self.assertIn(f"No {what} saved yet.", paragraphs)

    def test_a_missing_languages_table_is_stated_not_hidden(self):
        data = dict(SAMPLE, languages=[], languages_unavailable="Languages could not be included: the database upgrade has not been applied.")
        paragraphs, _ = text_of(build_profile_copy_bytes(data))
        self.assertIn("Languages could not be included: the database upgrade has not been applied.", paragraphs)

    def test_pretty_date(self):
        self.assertEqual(pretty_date("2026-10-09T01:19:44.808945+00:00"), "Oct 9, 2026")
        self.assertEqual(pretty_date(None), "")
        self.assertEqual(pretty_date("sometime"), "sometime")


if __name__ == "__main__":
    unittest.main()
