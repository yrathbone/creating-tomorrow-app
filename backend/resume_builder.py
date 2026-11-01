"""
Builds .docx files in memory (returns bytes rather than writing to disk,
since this runs inside a web request).

build_resume_bytes() produces the main resume, matching the Creating
Tomorrow template layout:

  NAME (bold, centered)
  Location | Phone | Email | LinkedIn   (centered)
  [shaded bar] PROFESSIONAL SUMMARY
  Summary paragraph(s)
  [shaded bar] CORE SKILLS & EXPERTISE
  Skills bullets, two columns by default (or one in ATS mode)
  [shaded bar] PROFESSIONAL EXPERIENCE
  Job title (bold) / Company, Location - Dates
  Bulleted achievements
  ... (repeat per job)
  [shaded bar] EDUCATION
  Bulleted degree/school lines

build_resume_bytes(data, ats_mode, template) takes a layout name: "classic" (the layout above,
unchanged), "modern" or "traditional". Unknown names quietly mean classic. See TEMPLATES.

build_match_recap_bytes() produces a one-page downloadable recap of a
Right Fit comparison (match level, rationale, strengths, gaps, flags,
growth suggestions) using the same visual style.

build_resume_bytes() also recognizes a few optional keys, harmless to any
tool that doesn't set them (absent for them, so nothing changes):
  data["headline"]           - centered bold positioning line under contact
  data["skills_heading"]     - overrides the "CORE SKILLS & EXPERTISE"
                                heading text (Elevate/Refine use "CORE
                                EXPERTISE")
  data["certifications"]     - bulleted section rendered after Education,
                                only if non-empty (used by Elevate)
  data["lang"]              - the document language tag, e.g. "es-US" (default "en-US")
  data["additional_sections"] - a list of {"heading": str, "items": [str]}
                                sections rendered after Education/
                                Certifications, one heading bar + bulleted
                                items per entry, only for entries with a
                                non-empty items list (used by Refine to
                                preserve things like Awards & Recognition,
                                Licenses, Languages, Military Service,
                                Professional Affiliations, Publications,
                                Patents, or Security Clearances found in the
                                source resume)
"""
import copy
import io
import re
from datetime import datetime, timezone

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEADING_FILL = "CADCF2"
FONT_NAME = "Calibri"
FONT_SIZE = Pt(11)
PAGE_MARGIN = Inches(0.7)

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT

# Resume layouts. The words are identical in every layout; only the look changes.
#   classic     - the original Creating Tomorrow look: centered, shaded section bars (unchanged).
#   modern      - left-aligned, a large navy name, section titles with a thin navy line under them.
#   traditional - serif font, centered name, section titles with a thin black line under them.
# "heading" is "bar" (shaded paragraph) or "rule" (a line under the title).
TEMPLATES = {
    "classic": {
        "font": "Calibri", "size": 11, "name_size": 11, "name_color": None,
        "contact_size": 11, "contact_color": None, "align": CENTER,
        "heading": "bar", "heading_fill": "CADCF2", "heading_color": None,
        "heading_align": CENTER, "rule_color": None, "heading_before": None,
    },
    "modern": {
        "font": "Calibri", "size": 10.5, "name_size": 22, "name_color": "1F3A8A",
        "contact_size": 10, "contact_color": "4B5565", "align": LEFT,
        "heading": "rule", "heading_fill": None, "heading_color": "1F3A8A",
        "heading_align": LEFT, "rule_color": "1F3A8A", "heading_before": 10,
    },
    "traditional": {
        "font": "Times New Roman", "size": 11, "name_size": 18, "name_color": None,
        "contact_size": 10.5, "contact_color": None, "align": CENTER,
        "heading": "rule", "heading_fill": None, "heading_color": None,
        "heading_align": LEFT, "rule_color": "000000", "heading_before": 8,
    },
}
DEFAULT_TEMPLATE = "classic"


def resolve_template(name) -> dict:
    """The layout for a name from the request; anything unknown quietly means Classic."""
    return TEMPLATES.get(name if isinstance(name, str) else "", TEMPLATES[DEFAULT_TEMPLATE])


def shade_paragraph(paragraph, fill_hex):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    pPr.insert_element_before(shd, *_AFTER_PBDR[1:])  # shading belongs before alignment and spacing (it was appended last)


# Where a bottom border belongs among a paragraph's properties (Word wants them in this order).
_AFTER_PBDR = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
    "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
    "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
    "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def underline_paragraph(paragraph, color_hex):
    """A thin line under the paragraph (used by the Modern and Traditional section titles)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.insert_element_before(pBdr, *_AFTER_PBDR)


def set_columns(section, num_cols):
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num_cols))
    cols.set(qn("w:space"), "720")


def style_run(run, bold=False, tpl=None, size=None, color=None):
    tpl = tpl or TEMPLATES[DEFAULT_TEMPLATE]
    run.font.name = tpl["font"]
    run.font.size = Pt(size if size is not None else tpl["size"])
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), tpl["font"])


def add_heading_bar(doc, text, tpl=None):
    tpl = tpl or TEMPLATES[DEFAULT_TEMPLATE]
    p = doc.add_paragraph(style="Heading 1")  # a real Word heading (Navigation Pane, screen readers); prepare_heading_style() keeps the look
    p.alignment = tpl["heading_align"]
    run = p.add_run(text)
    style_run(run, bold=True, tpl=tpl, color=tpl["heading_color"])
    if tpl["heading"] == "bar":
        shade_paragraph(p, tpl["heading_fill"])
    else:
        p.paragraph_format.space_before = Pt(tpl["heading_before"] or 8)
        p.paragraph_format.space_after = Pt(3)
        underline_paragraph(p, tpl["rule_color"])
    return p


def add_bullet(doc, text, tpl=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(clean_text(text))
    style_run(run, tpl=tpl)
    return p


def add_skill_line(doc, text, tpl=None):
    """A compact skills line like "Payments & Integration: API Connectivity, ISO 20022"
    - the part before the first colon (if any) is the bold group label."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    text = clean_text(text)
    label, sep, rest = text.partition(":")
    if sep and len(label) <= 40:
        style_run(p.add_run(label + ":"), bold=True, tpl=tpl)
        style_run(p.add_run(rest), tpl=tpl)
    else:
        style_run(p.add_run(text), tpl=tpl)
    return p


def set_page_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = PAGE_MARGIN
    section.right_margin = PAGE_MARGIN
    section.top_margin = PAGE_MARGIN
    section.bottom_margin = PAGE_MARGIN


# ---- text that arrives from outside (an AI answer, a saved resume, a pasted line) can be messy; make it safe -----------

DEFAULT_LANG = "en-US"
_LANG = re.compile(r"^[A-Za-z]{2,3}(-[A-Za-z0-9]{2,8}){0,2}$")
_LINE_BREAKS = re.compile(r"[\x0b\x0c\u2028\u2029]|\r\n?")  # Word's soft line break, form feed, Unicode separators, CR/CRLF -> a plain line break
_UNUSABLE = re.compile(r"[\x00-\x08\x0e-\x1f\x7f\ufffe\uffff\ud800-\udfff]")  # not allowed in a Word file: removed instead of crashing the build
_INVISIBLE = re.compile(r"[\u200b\u2060\ufeff\u202a-\u202e]")  # zero-width space, word joiner, byte-order mark, bidi overrides (ZWJ/ZWNJ are kept: some scripts need them)
_LEADING_BULLET = re.compile(r"^\s*[\u2022\u25cf\u25aa\u25e6\u25a0\u2219\u00b7\u2023\u2043]+\s*")  # a typed bullet in front of a list item would give two bullets


def clean_text(value) -> str:
    """Any value as safe text for a Word file: None is empty, numbers become text, unusable control characters are dropped,
    vertical tabs and similar become line breaks, invisible characters are removed."""
    if value is None:
        return ""
    text = value if isinstance(value, str) else str(value)
    text = _LINE_BREAKS.sub("\n", text)
    text = _UNUSABLE.sub("", text)
    return _INVISIBLE.sub("", text)


def clean_line(value) -> str:
    """Single-line text (a name, a heading, a job title): line breaks become spaces."""
    return re.sub(r"\s+", " ", clean_text(value)).strip()


def clean_items(value) -> list:
    """A list of text items. A lone string is ONE item (never split into characters); None entries and empty items are dropped;
    a bullet character typed at the start of an item is removed because the list adds its own."""
    if value is None:
        return []
    values = value if isinstance(value, (list, tuple)) else [value]
    items = []
    for v in values:
        text = _LEADING_BULLET.sub("", clean_text(v)).strip()
        if text:
            items.append(text)
    return items


def clean_structure(value):
    """Every text inside a nested dict/list made safe (used for the recap documents)."""
    if isinstance(value, dict):
        return {k: clean_structure(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean_structure(v) for v in value]
    if isinstance(value, str):
        return clean_text(value)
    return value


def normalize_resume(data) -> dict:
    """The resume as the builder needs it: only text, lists of text, and the known keys. Missing parts are empty (the
    sections are then simply left out); nothing here changes the words of well-formed data."""
    d = data if isinstance(data, dict) else {}
    summary = d.get("summary")
    summary_items = [clean_text(x).strip() for x in (summary if isinstance(summary, (list, tuple)) else [summary])]
    jobs = []
    for entry in d.get("experience") or []:
        if isinstance(entry, str):
            entry = {"title": entry}
        if not isinstance(entry, dict):
            continue
        job = {"title": clean_line(entry.get("title")), "subtitle": clean_line(entry.get("subtitle")), "bullets": clean_items(entry.get("bullets"))}
        if job["title"] or job["subtitle"] or job["bullets"]:
            jobs.append(job)
    extra = []
    for section in d.get("additional_sections") or []:
        if isinstance(section, dict) and clean_items(section.get("items")):
            extra.append({"heading": clean_line(section.get("heading")), "items": clean_items(section.get("items"))})
    lang = clean_line(d.get("lang") or d.get("language"))
    return {
        "name": clean_line(d.get("name")),
        "contact": clean_line(d.get("contact")),
        "headline": clean_line(d.get("headline")),
        "summary": [x for x in summary_items if x],
        "skills": clean_items(d.get("skills")),
        "skills_heading": clean_line(d.get("skills_heading")) or "CORE SKILLS & EXPERTISE",
        "skills_style": "grouped" if d.get("skills_style") == "grouped" else "",
        "experience": jobs,
        "education": clean_items(d.get("education")),
        "certifications": clean_items(d.get("certifications")),
        "additional_sections": extra,
        "lang": lang if _LANG.match(lang) else DEFAULT_LANG,
    }


# ---- document structure: heading style, properties, language, column sections ------------------------------------------

def prepare_heading_style(doc, tpl):
    """Word's built-in "Heading 1" is used for every section title so the file has real headings. Its stock look (blue, 14 pt,
    24 pt of space above) is replaced here by the layout's own font, size, colour and spacing, so nothing looks different."""
    style = doc.styles["Heading 1"]
    rPr = style.element.get_or_add_rPr()
    for old in rPr.findall(qn("w:rFonts")):
        rPr.remove(old)
    fonts = OxmlElement("w:rFonts")
    for attribute in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn("w:" + attribute), tpl["font"])
    rPr.insert(0, fonts)
    style.font.size = Pt(tpl["size"])
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    color = rPr.find(qn("w:color"))
    color.set(qn("w:val"), "auto")  # "automatic" colour, like ordinary text
    for attribute in ("themeColor", "themeShade", "themeTint"):
        color.attrib.pop(qn("w:" + attribute), None)
    spacing = style.paragraph_format
    spacing.space_before = Pt(0)
    spacing.space_after = Pt(10)
    spacing.keep_with_next = True  # a section title is never left alone at the bottom of a page
    spacing.keep_together = True


def set_language(doc, lang):
    """The document's proofing and reading language (it was fixed at English)."""
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    lang_element = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:lang")) if defaults is not None else None
    if lang_element is None:
        return
    lang_element.set(qn("w:val"), lang)
    lang_element.set(qn("w:eastAsia"), lang)


def set_properties(doc, title, author="", lang=DEFAULT_LANG):
    """File properties a person (or a recruiter's software) can read: a real title, the person as author, today's date.
    The stock template said the author was "python-docx" and the date was 2013."""
    now = datetime.now(timezone.utc).replace(microsecond=0, tzinfo=None)
    core = doc.core_properties
    core.title = title
    core.author = author
    core.last_modified_by = ""
    core.comments = ""
    core.keywords = ""
    core.subject = ""
    core.category = ""
    core.language = lang
    core.created = now
    core.modified = now
    core.revision = 1


def end_section_after(doc, paragraph, columns=None, continuous=False):
    """Close a section at `paragraph` itself (Word's own way), instead of adding an empty paragraph just to carry the break."""
    sect = copy.deepcopy(doc.sections[0]._sectPr)
    for tag in ("w:type", "w:cols"):
        for old in sect.findall(qn(tag)):
            sect.remove(old)
    if continuous:
        kind = OxmlElement("w:type")
        kind.set(qn("w:val"), "continuous")
        sect.insert(0, kind)
    if columns:
        cols = OxmlElement("w:cols")
        cols.set(qn("w:num"), str(columns))
        cols.set(qn("w:space"), "720")
        grid = sect.find(qn("w:docGrid"))
        if grid is not None:
            grid.addprevious(cols)
        else:
            sect.append(cols)
    paragraph._p.get_or_add_pPr()._insert_sectPr(sect)


def build_resume_bytes(data: dict, ats_mode: bool = False, template: str = DEFAULT_TEMPLATE) -> bytes:
    tpl = resolve_template(template)
    resume = normalize_resume(data)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = tpl["font"]
    normal.font.size = Pt(tpl["size"])
    prepare_heading_style(doc, tpl)
    set_language(doc, resume["lang"])
    set_page_geometry(doc.sections[0])
    set_properties(doc, f"{resume['name']} \u2014 R\u00e9sum\u00e9" if resume["name"] else "R\u00e9sum\u00e9", resume["name"], resume["lang"])

    # a rule-style layout used to get a tiny empty paragraph after the skills list as breathing room; the same space now
    # comes from the next section title instead (no empty paragraph in the file)
    space_owed_after_skills = [0]

    def section_title(text, space_before=None):
        heading = add_heading_bar(doc, text, tpl)
        if space_before is not None:
            heading.paragraph_format.space_before = Pt(space_before)
        if space_owed_after_skills[0]:
            current = heading.paragraph_format.space_before
            heading.paragraph_format.space_before = Pt((current.pt if current is not None else 0) + space_owed_after_skills[0])
            space_owed_after_skills[0] = 0
        return heading

    p = doc.add_paragraph()
    p.alignment = tpl["align"]
    name_run = p.add_run(resume["name"])
    style_run(name_run, bold=True, tpl=tpl, size=tpl["name_size"], color=tpl["name_color"])
    p.add_run().add_break()
    contact_run = p.add_run(resume["contact"])
    style_run(contact_run, tpl=tpl, size=tpl["contact_size"], color=tpl["contact_color"])

    if resume["headline"]:
        p = doc.add_paragraph()
        p.alignment = tpl["align"]
        p.paragraph_format.space_before = Pt(4)
        headline_run = p.add_run(resume["headline"])
        style_run(headline_run, bold=True, tpl=tpl)

    if resume["summary"]:
        section_title("PROFESSIONAL SUMMARY")
        p = doc.add_paragraph()
        for i, para_text in enumerate(resume["summary"]):
            if i > 0:
                p.add_run().add_break()
                p.add_run().add_break()
            run = p.add_run(para_text)
            style_run(run, tpl=tpl)

    # Sections, skills, experience and education only appear when there is real content (no orphaned heading bars).
    skills = resume["skills"]
    if skills:
        heading = section_title(resume["skills_heading"])
        grouped = resume["skills_style"] == "grouped"
        # "grouped" = a few compact "Label: skill, skill" lines (single column) instead of one bullet per skill in two columns;
        # ATS mode is single column too. Only the two-column list needs section breaks, and they sit on real paragraphs.
        two_columns = not (ats_mode or grouped)
        if two_columns:
            end_section_after(doc, heading)
        last = None
        for skill in skills:
            last = (add_skill_line if grouped else add_bullet)(doc, skill, tpl)
        if two_columns:
            end_section_after(doc, last, columns=2, continuous=True)
            doc.sections[-1].start_type = WD_SECTION.CONTINUOUS  # everything after the list is a normal single-column section again
        if tpl["heading"] == "rule":
            space_owed_after_skills[0] = 6  # the next section title carries the breathing room the old empty paragraph gave

    if resume["experience"]:
        section_title("PROFESSIONAL EXPERIENCE")
        for i, job in enumerate(resume["experience"]):
            p = doc.add_paragraph()
            if i > 0:
                p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True  # a job's title line stays with its first bullet
            p.paragraph_format.keep_together = True
            if job["title"]:
                style_run(p.add_run(job["title"]), bold=True, tpl=tpl)
            if job["title"] and job["subtitle"]:
                p.add_run().add_break()
            if job["subtitle"]:
                style_run(p.add_run(job["subtitle"]), tpl=tpl)
            for bullet in job["bullets"]:
                add_bullet(doc, bullet, tpl)

    if resume["education"]:
        section_title("EDUCATION", space_before=6)
        for entry in resume["education"]:
            add_bullet(doc, entry, tpl)

    if resume["certifications"]:
        section_title("CERTIFICATIONS", space_before=6)
        for entry in resume["certifications"]:
            add_bullet(doc, entry, tpl)

    for section in resume["additional_sections"]:
        section_title(section["heading"], space_before=6)
        for entry in section["items"]:
            add_bullet(doc, entry, tpl)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_match_recap_bytes(match_report: dict, candidate_name: str) -> bytes:
    """One-page .docx recap of a Right Fit comparison: match level, the
    rationale, strengths, gaps, any "same word, different job" flags, and
    growth suggestions - a downloadable summary of what's shown on screen.
    """
    match_report = clean_structure(match_report) if isinstance(match_report, dict) else {}
    candidate_name = clean_line(candidate_name)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    prepare_heading_style(doc, TEMPLATES[DEFAULT_TEMPLATE])
    set_properties(doc, f"{candidate_name} — Right Fit Recap" if candidate_name else "Right Fit Recap", candidate_name)

    set_page_geometry(doc.sections[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = p.add_run(f"{candidate_name} — Right Fit Recap")
    style_run(name_run, bold=True)

    add_heading_bar(doc, "MATCH LEVEL")
    p = doc.add_paragraph()
    level_run = p.add_run(match_report.get("match_level", ""))
    style_run(level_run, bold=True)
    if match_report.get("match_rationale"):
        p = doc.add_paragraph()
        style_run(p.add_run(match_report["match_rationale"]))

    strengths = match_report.get("strengths") or []
    if strengths:
        add_heading_bar(doc, "WHAT'S ALREADY STRONG")
        for item in strengths:
            add_bullet(doc, item)

    gaps = match_report.get("required_qualification_gaps") or []
    if gaps:
        add_heading_bar(doc, "GAPS AGAINST WHAT THIS ROLE USUALLY REQUIRES")
        for gap in gaps:
            text = f"[{gap.get('status', '')}] {gap.get('requirement', '')} — {gap.get('explanation', '')}"
            add_bullet(doc, text)

    flags = match_report.get("same_word_different_job_flags") or []
    if flags:
        add_heading_bar(doc, '"SAME WORD, DIFFERENT JOB" FLAGS')
        for flag in flags:
            text = (
                f"“{flag.get('term', '')}” — on the resume: {flag.get('resume_meaning', '')}. "
                f"In the posting: {flag.get('posting_meaning', '')}. {flag.get('why_it_matters', '')}"
            )
            add_bullet(doc, text)

    growth = match_report.get("growth_suggestions") or []
    if growth:
        add_heading_bar(doc, "HOW TO GENUINELY GROW TOWARD THIS ROLE")
        for item in growth:
            add_bullet(doc, item)

    if match_report.get("note_on_better_fit_roles"):
        add_heading_bar(doc, "A NOTE ON FIT")
        p = doc.add_paragraph()
        style_run(p.add_run(match_report["note_on_better_fit_roles"]))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_profile_review_recap_bytes(review: dict) -> bytes:
    """One-page .docx summary of a Spotlight review: signature strengths,
    the concrete suggested headline/About rewrites (the parts someone can
    actually copy and use), a consolidated list of ways to strengthen the
    profile, and - when there was enough material - a sample of what an
    updated profile could look like altogether, at the end. Deliberately
    tighter than the full on-screen section-by-section breakdown, so the
    download is something worth acting on rather than a re-read of the
    whole page.
    """
    review = clean_structure(review) if isinstance(review, dict) else {}
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = FONT_SIZE
    prepare_heading_style(doc, TEMPLATES[DEFAULT_TEMPLATE])
    set_properties(doc, "Spotlight — LinkedIn Profile Review")

    set_page_geometry(doc.sections[0])

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = p.add_run("Spotlight — LinkedIn Profile Review")
    style_run(name_run, bold=True)

    if review.get("insufficient_information"):
        add_heading_bar(doc, "MORE INFORMATION NEEDED")
        p = doc.add_paragraph()
        style_run(p.add_run(
            review.get("insufficient_information_message")
            or "There wasn't enough profile information supplied for a useful review."
        ))
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    strengths = review.get("strengths") or []
    if strengths:
        add_heading_bar(doc, "SIGNATURE STRENGTHS")
        for item in strengths:
            p = doc.add_paragraph()
            title_run = p.add_run(item.get("title", ""))
            style_run(title_run, bold=True)
            if item.get("evidence"):
                ev = doc.add_paragraph()
                style_run(ev.add_run(item["evidence"]))

    headline = review.get("headline") or {}
    if headline.get("supplied") and headline.get("suggested"):
        add_heading_bar(doc, "SUGGESTED HEADLINE")
        p = doc.add_paragraph()
        style_run(p.add_run(headline["suggested"]))

    about = review.get("about") or {}
    if about.get("supplied") and about.get("suggested_revision"):
        add_heading_bar(doc, "SUGGESTED ABOUT SECTION")
        p = doc.add_paragraph()
        style_run(p.add_run(about["suggested_revision"]))

    ways_to_strengthen = []
    if headline.get("supplied"):
        ways_to_strengthen.extend(headline.get("could_be_clearer") or [])
    if about.get("supplied"):
        ways_to_strengthen.extend(about.get("could_be_clearer") or [])
    experience = review.get("experience") or {}
    if experience.get("supplied"):
        ways_to_strengthen.extend(experience.get("suggested_improvements") or [])
    skills = review.get("skills") or {}
    ways_to_strengthen.extend(skills.get("needs_more_information") or [])

    if ways_to_strengthen:
        add_heading_bar(doc, "WAYS TO STRENGTHEN THIS PROFILE")
        for item in ways_to_strengthen:
            add_bullet(doc, item)

    sample = review.get("suggested_full_profile")
    if sample:
        add_heading_bar(doc, "SAMPLE UPDATED PROFILE")
        p = doc.add_paragraph()
        style_run(p.add_run(
            "One way to put these suggestions together - feel free to use, adapt, or ignore any of it."
        ))
        if sample.get("headline"):
            p = doc.add_paragraph()
            style_run(p.add_run("Headline: "), bold=True)
            style_run(p.add_run(sample["headline"]))
        if sample.get("about"):
            hp = doc.add_paragraph()
            style_run(hp.add_run("About"), bold=True)
            p = doc.add_paragraph()
            style_run(p.add_run(sample["about"]))
        for role in sample.get("experience") or []:
            title = role.get("title", "")
            org = role.get("organization", "")
            p = doc.add_paragraph()
            style_run(p.add_run(f"{title} — {org}" if org else title), bold=True)
            for bullet in role.get("bullets") or []:
                add_bullet(doc, bullet)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
