"""A readable copy of a person's saved Career Profile, as a Word document.

The data-file backup (/api/career/export) is complete but is a technical format. This turns the very same snapshot
into something a person can read, print or keep: roles, education, certifications, skills, languages, applications and a
short history of scans and resumes. It only formats what is already saved. Nothing is added, rewritten or invented.
"""
import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INK = RGBColor(0x14, 0x21, 0x3D)
BLUE = RGBColor(0x0A, 0x35, 0xA8)
MUTED = RGBColor(0x5A, 0x64, 0x78)

SCAN_LABELS = {"job_comparison": "Job comparison", "skill_scan": "Skill scan"}


def pretty_date(value) -> str:
    """'2026-10-09T01:19:44+00:00' -> 'Oct 9, 2026'. Anything unreadable is shown as it was saved."""
    if not value:
        return ""
    try:
        d = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return f"{d.strftime('%b')} {d.day}, {d.year}"
    except ValueError:
        return str(value)


def _lines(text) -> list[str]:
    return [" ".join(line.split()) for line in str(text or "").splitlines() if line.strip()]


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = BLUE
    pPr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "C3CDE6")):
        bottom.set(qn(f"w:{k}"), v)
    border.append(bottom)
    pPr.append(border)


def _small(doc, text, italic=False, color=MUTED, after=2, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    r.italic = italic
    r.font.size = Pt(10)
    r.font.color.rgb = color
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.add_run(text)


def _none_yet(doc, what):
    _small(doc, f"No {what} saved yet.", italic=True)


def _join(parts, sep=" · "):
    return sep.join(str(p) for p in parts if p)


def _span(start, end) -> str:
    if start and end:
        return f"{start} to {end}"
    return start or end or ""


def build_profile_copy_bytes(data: dict, name: str | None = None, contact: str | None = None) -> bytes:
    """`data` is the dictionary the export route produces (applications, experiences, education, certifications,
    skills, scan_history, languages, resume_versions, exported_at)."""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(4)

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(0)
    r = t.add_run((name or "").strip() or "My Career Profile")
    r.bold = True
    r.font.size = Pt(24)
    r.font.name = "Georgia"
    r.font.color.rgb = INK
    if (name or "").strip():
        _small(doc, "Career Profile", color=BLUE, after=0)
    if (contact or "").strip():
        _small(doc, contact.strip(), after=2)
    _small(
        doc,
        f"A readable copy of what is saved in your Creating Tomorrow Career Profile, as of {pretty_date(data.get('exported_at'))}. "
        "It shows only what is saved. For a complete copy that can be restored or moved, also keep the full data backup.",
        italic=True, after=6,
    )

    experiences = data.get("experiences") or []
    _heading(doc, "Work experience")
    if not experiences:
        _none_yet(doc, "roles")
    for e in experiences:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(_join([e.get("title"), e.get("organization")], " at ") or "Role")
        r.bold = True
        r.font.size = Pt(11.5)
        meta = _join([e.get("location"), _span(e.get("start_date"), e.get("end_date")), "" if e.get("verified") else "not yet confirmed by you"])
        if meta:
            _small(doc, meta, keep=True)
        for line in _lines(e.get("description")):
            _bullet(doc, line)

    education = data.get("education") or []
    _heading(doc, "Education")
    if not education:
        _none_yet(doc, "education")
    for ed in education:
        degree = _join([ed.get("degree"), ed.get("field_of_study")], ", ")
        _bullet(doc, _join([degree, ed.get("institution"), ed.get("graduation_date")], " · "))

    certs = data.get("certifications") or []
    _heading(doc, "Certifications")
    if not certs:
        _none_yet(doc, "certifications")
    for c in certs:
        _bullet(doc, _join([c.get("name"), c.get("issuer"), c.get("date")], " · "))

    skills = sorted({(s.get("name") or "").strip() for s in (data.get("skills") or []) if (s.get("name") or "").strip()}, key=str.lower)
    _heading(doc, f"Skills ({len(skills)})" if skills else "Skills")
    if skills:
        p = doc.add_paragraph()
        p.add_run(", ".join(skills))
    else:
        _none_yet(doc, "skills")

    languages = data.get("languages") or []
    if languages:
        _heading(doc, "Languages")
        for lang in languages:
            _bullet(doc, _join([lang.get("name"), lang.get("proficiency")], " · "))
    elif data.get("languages_unavailable"):
        _heading(doc, "Languages")
        _small(doc, data["languages_unavailable"], italic=True)

    applications = data.get("applications") or []
    _heading(doc, "Applications you are tracking")
    if not applications:
        _none_yet(doc, "applications")
    else:
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Grid Accent 1"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, text in zip(table.rows[0].cells, ("Job", "Company", "Applied", "Status", "Notes")):
            cell.text = text
        for a in applications:
            cells = table.add_row().cells
            for cell, text in zip(cells, (a.get("job_title"), a.get("company"), a.get("applied_on"), a.get("status"), a.get("notes"))):
                cell.text = str(text or "")
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(9.5)

    scans = data.get("scan_history") or []
    titles_by_scan = {s.get("id"): s.get("job_title") for s in scans}
    _heading(doc, "Scans you have run")
    if not scans:
        _none_yet(doc, "scans")
    for s in scans:
        label = SCAN_LABELS.get(s.get("scan_type"), (s.get("scan_type") or "Scan").replace("_", " ").capitalize())
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(_join([label, s.get("job_title")], ": "))
        r.bold = True
        _small(doc, pretty_date(s.get("created_at")), keep=bool(s.get("summary_text")))
        for line in _lines(s.get("summary_text")):
            _bullet(doc, line)

    resumes = data.get("resume_versions") or []
    _heading(doc, "Resumes you have built")
    if not resumes:
        _none_yet(doc, "resumes")
    else:
        _small(doc, "Download any of them again, in the layout you prefer, from Resumes Built on your dashboard.", italic=True, after=4)
    for v in resumes:
        target = titles_by_scan.get(v.get("scan_history_id"))
        who = ((v.get("resume_data") or {}).get("name") or "").strip()
        _bullet(doc, _join([pretty_date(v.get("created_at")), f"for {target}" if target else "general resume", who]))

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
